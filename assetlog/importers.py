"""Statement-file imports: KIS [7451] 거래내역 (xlsx/csv/html-xls), 토스증권 거래내역서 PDF, generic CSV.

Each row becomes a ledger transaction: KRW, signed cash effect (+ in, - out). The memo is the source's
거래종류/거래구분 verbatim, never the 종목명 (a buy of an ETF named '…배당…' would classify as DIVIDEND).
Dividends are passed as net cash received: gross - withholding on KRW rows; Toss USD rows are already net.

A file is the booked record: it replaces the estimates (infer.py flows, computed `_api` dividends) that would count its
money twice, including ones the user edited in 검토 (listed in the result as replaced_edited). A complete statement
(COMPLETE) also records the days it covers (db coverage), so sync and infer.py don't estimate them again; it replaces
the flow estimates whose window lies inside that coverage and the computed dividends inside its days. Any file
replaces the flow estimates whose window holds a deposit/withdrawal it booked, and a partial one (csv) the computed
dividends on its dividend days, never kinds it doesn't contain. An estimate whose window starts before the file's
first day also stands for days the file doesn't cover, so the import leaves it.

Unverified against a live account (no real file was available; layouts follow docs/contracts/imports.json):
- KIS: columns are found by header name (거래일, 거래종류|거래상세유형, 거래금액, 정산금액, 수수료, 거래세, 세금,
  환율, 잔액|예수금; 입출금 tab: 입금액, 출금액). Only the 전체거래내역 CSV header was seen in real files; the
  HTS xlsx/xls and 입출금거래내역 download headers are assumed. Missing columns raise with the headers found.
- KIS: a 거래종류 cell without Hangul is taken as the code sub-cell and the text is read from the next cell, when no
  header names that column (so never from 종목명).
- KIS: rows with a nonzero 환율 are assumed to hold 거래금액/세금 in that foreign currency and are multiplied by
  환율. Until a real export confirms that, such rows go to review (UNKNOWN) unless they classify as INTERNAL, so
  the guess never reaches deposits/dividends unchecked. The overseas dividend/withholding 거래종류 strings are unknown.
- KIS: dividends = 거래금액 - 세금. No taxed ETF분배금입금 was ever seen, so the row is imported as UNKNOWN
  (review queue) when the file has no 세금 column or 정산금액 disagrees with 거래금액 - 세금.
- KIS: 연금납입금전환 is skipped as a non-cash marker. Importing both the 전체거래내역 and the 입출금거래내역 of
  one account dedupes only if 잔액 (전체) equals 예수금 (입출금) after each row; import one tab per account.
- Toss: checked against a real 거래내역서 (2026-01-01..09-25, 391 rows). 거래구분 is the row's first word, or one of
  TOSS_MULTI; the rest is the 종목명, kept in the memo but never used to classify. 환율 is blank on plain KRW rows.
  The cash effect is 정산금액 (net of fee and tax) with its sign from 거래구분; USD rows carry it in KRW on the first
  line and in $ on the second. Every row must move the running 잔액 of its currency by exactly that much (KRW ±1,
  USD ±$0.01), which catches a misread column or a tax posted somewhere else.
- Toss: TOSS_KINDS strings other than 구매/판매/대여료 입금 come from parser maps, not real rows.
"""
import csv
import hashlib
import io
import itertools
import json
import re
import shutil
import subprocess
import zipfile
from datetime import date, timedelta

from . import db, dcmodel, infer

FORMATS = {
    "kis_xlsx": "한국투자증권 계좌별 거래내역 (전체/입출금, xlsx·csv·xls)",
    "toss_pdf": "토스증권 거래내역서 (PDF)",
    "csv": "기타 CSV (날짜·금액·적요)",
}
COMPLETE = ("toss_pdf", "kis_xlsx")  # every transaction of their days; a csv may be a partial transcription
FIRST_DAY = date(1990, 1, 1)  # dates accepted: as manual entry (web._past), up to tomorrow
MAX_WON = 1e13      # as manual entry; also rejects inf/nan
MAX_ROWS, MAX_COLS = 50_000, 64  # a statement is far smaller; a crafted file (far cell, huge span) must not fill RAM
MAX_UNZIPPED = 200 * 2**20  # xlsx
PDF_TIMEOUT, MAX_TEXT = 60, 50 * 2**20  # a 26-page 거래내역서: well under a second, ~100KB of text

# First match sets the direction; anything else is money in.
SIGNS = (("입금취소", -1), ("출금취소", 1), ("환급", 1), ("환불", 1), ("출금", -1), ("출고", -1), ("매수", -1),
         ("구매", -1), ("세금", -1), ("세액", -1), ("배당세", -1), ("현지세", -1), ("추징", -1), ("수수료", -1),
         ("청약", -1))
INCOME = ("배당", "분배금", "이용료", "이자", "대여료")
KEEP_ZERO = ("입고", "출고")  # securities moved: no cash, but the user must put a value on them
NON_CASH = ("연금납입금전환",)  # KIS: 잔액 unchanged, mirrors a same-day 타사이체입금 of the same amount
# Toss strings classify() gets wrong; 구매/판매 would otherwise all land in the review queue.
TOSS_KINDS = {"구매": "INTERNAL", "판매": "INTERNAL", "대여료 입금": "INTEREST", "액면분할단주대금입금": "INTERNAL",
              "외국납부세액출금": "DIVIDEND", "외화세금환급": "DIVIDEND"}
TOSS_MULTI = ("이벤트 입금", "대여료 입금", "대차거래 대여출고", "대차거래 상환입고")  # 거래구분 with a space
HANGUL = re.compile("[가-힣]")
NUM = r"-?[\d,]+(?:\.\d+)?"
TOSS_DATE = re.compile(r"\s*\d{4}\.\d\d\.\d\d\s")
USD_VALUES = re.compile(rf"\(\$\s*({NUM})\)")
# Real 거래내역서 (2026-09): 환율 is blank on plain KRW rows; USD rows repeat their money in ($ …) on the next line.
KRW_TAIL = ("거래수량", "거래대금", "정산금액", "단가", "수수료", "거래세", "제세금", "변제/연체합", "잔고", "잔액")
USD_TAIL = ("거래수량", "거래대금", "정산금액", "단가", "수수료", "제세금", "변제/연체합", "잔고", "잔액")
USD_LINE = ("거래대금", "정산금액", "단가", "수수료", "제세금", "변제/연체합", "잔액")  # the ($ …) values
CSV_COLS = {
    "date": ("날짜", "거래일자", "거래일", "일자", "date"),
    "amount": ("금액", "amount"),  # signed; an unsigned 거래금액 (KIS) must go through kis_xlsx
    "in": ("입금액", "입금"), "out": ("출금액", "출금"),
    "memo": ("적요", "내용", "memo", "거래구분", "거래종류", "description"),
    "bal": ("잔액", "balance"),
}


def import_bytes(conn, account_id, fmt, data: bytes, filename=""):
    if fmt not in PARSERS:
        raise ValueError(f"unknown format {fmt!r}; expected one of {sorted(FORMATS)}")
    rows = list(PARSERS[fmt](data))  # parse everything first: a bad file writes nothing
    dated = sorted(d.isoformat() for d, *_ in rows if d)
    if not dated:
        raise ValueError(f"{filename or fmt}: no dated transaction rows found")
    for d, amount, memo, *_ in rows:
        if d and not FIRST_DAY <= d <= date.today() + timedelta(days=1):
            raise ValueError(f"{filename or fmt}: row {d} {memo!r}: date outside {FIRST_DAY}..today")
        if amount is not None and not abs(amount) < MAX_WON:
            raise ValueError(f"{filename or fmt}: row {d} {memo!r}: amount {amount} out of range")
    out, seen, booked, dc = {"imported": 0, "unknown": 0, "skipped": 0}, {}, {}, 0
    with conn:
        for d, amount, memo, bal, raw, kind in rows:
            if (amount is None or any(k in memo for k in NON_CASH)
                    or (not amount and not any(k in memo for k in KEEP_ZERO))):
                out["skipped"] += 1
                continue
            amount = round(amount, 2)
            kind = kind or db.classify(memo, amount)
            # Identical rows on one day get #1, #2…: stable as long as each export covers whole days.
            base = f"{d}|{amount}|{memo}|{bal}"
            seen[base] = n = seen.get(base, 0) + 1
            key = hashlib.sha1(f"{base}|{n}".encode()).hexdigest()[:16]
            db.upsert_transaction(conn, account_id, d, amount, memo, fmt, key, kind, json.dumps(raw, ensure_ascii=False))
            booked.setdefault(kind, set()).add(d.isoformat())
            out["imported"] += 1
            out["unknown"] += kind == "UNKNOWN"
            if kind in ("DEPOSIT", "DIVIDEND"):
                before = conn.total_changes
                dcmodel.yield_month(conn, account_id, d, kind)
                dc += conn.total_changes - before
        days = dated if fmt in COMPLETE else sorted(set().union(*booked.values()))
        gone = _replace_estimates(conn, account_id, fmt, days[0], days[-1], booked) if days else []
        if gone or dc:
            out["replaced_estimates"] = len(gone) + dc
        edited = [{k: r[k] for k in ("date", "kind", "amount", "memo")} for r in gone if r["kind_locked"]]
        if edited:
            out["replaced_edited"] = edited
    return out


def _replace_estimates(conn, account_id, fmt, first, last, booked):
    """Delete the estimates this file's rows now book (see the module docstring); returns the deleted rows.
    booked: {kind: {ISO day}} of the rows imported; first, last: the file's first and last day."""
    if fmt in COMPLETE:
        conn.execute("INSERT OR IGNORE INTO coverage VALUES (?, ?, ?, ?)", (account_id, fmt, first, last))
    ranges = db.coverage(conn, account_id) if fmt in COMPLETE else []
    flows = booked.get("DEPOSIT", set()) | booked.get("WITHDRAW", set())
    gone = [r for r, lo, hi in infer.estimates(conn, account_id)
            if any(f <= lo and hi <= t for f, t in ranges) or (lo >= first and any(lo <= d <= hi for d in flows))]
    divs = (first, last) if fmt in COMPLETE else sorted(booked.get("DIVIDEND", ()))
    gone += conn.execute(
        "SELECT id, date, kind, amount, memo, source, kind_locked FROM transactions WHERE account_id = ? "
        "AND source LIKE '%\\_api' ESCAPE '\\' AND source_key LIKE 'div:%' AND "
        + ("date BETWEEN ? AND ?" if fmt in COMPLETE else f"date IN ({', '.join('?' * len(divs))})"),
        (account_id, *divs)).fetchall()
    conn.executemany("DELETE FROM transactions WHERE id = ?", [(r["id"],) for r in gone])
    return gone


def _sign(memo):
    return next((s for k, s in SIGNS if k in memo), 1)


def _income(memo):
    return _sign(memo) > 0 and any(k in memo for k in INCOME)


def _cash(memo, gross, fee=0.0, tax=0.0, ttax=0.0, withheld=True):
    """Signed cash effect. withheld=False: `gross` is already net of tax (Toss USD rows)."""
    s = _sign(memo)
    if _income(memo):
        return gross - tax if withheld else gross
    if s < 0 and ("매수" in memo or "구매" in memo):
        return -(gross + fee + ttax)
    if s > 0 and ("매도" in memo or "판매" in memo):
        return gross - fee - ttax - tax
    return s * (gross or tax)


def _num(s):
    t = str(s).replace(",", "").replace(" ", "")
    if t in ("", "-"):
        return 0.0
    try:
        v = float(t)
    except ValueError:
        raise ValueError(f"not a number: {s!r}") from None
    if not abs(v) < MAX_WON:  # inf, nan, 1e300
        raise ValueError(f"not a plausible number: {s!r}")
    return v


def _date(s):
    m = re.match(r"\s*(\d{4})[.\-/]?\s*(\d{1,2})[.\-/]?\s*(\d{1,2})(?!\d)", s or "")
    try:
        return date(*map(int, m.groups())) if m else None
    except ValueError:
        return None


def _norm(name):
    return re.sub(r"\s+|\(.*?\)", "", name or "").lower()


def _first(row):
    return row[0].strip() if row else ""


def _grid(data):
    """Rows of cell strings from xlsx, HTML-table .xls or CSV (utf-8 or cp949) bytes."""
    if data[:4] == b"PK\x03\x04":
        from openpyxl import load_workbook
        if sum(i.file_size for i in zipfile.ZipFile(io.BytesIO(data)).infolist()) > MAX_UNZIPPED:
            raise ValueError(f"xlsx unpacks to over {MAX_UNZIPPED >> 20}MB: not a 거래내역")
        ws = load_workbook(io.BytesIO(data), read_only=True, data_only=True).worksheets[0]
        return _bounded(ws.iter_rows(max_row=MAX_ROWS + 1, max_col=MAX_COLS, values_only=True))
    if data[:8] == b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1":
        raise ValueError("binary .xls (BIFF) can't be read without xlrd: open it in Excel/LibreOffice and "
                         "save it as .xlsx or .csv, then import that")
    try:
        text = data.decode("utf-8-sig")
    except UnicodeDecodeError:
        text = data.decode("cp949")
    if text.lstrip()[:1] == "<":
        return _bounded(_html_rows(text))
    return _bounded(csv.reader(io.StringIO(text)))


def _bounded(rows):
    """At most MAX_ROWS rows of MAX_COLS cell strings; a longer file fails rather than losing its tail."""
    out = [["" if v is None else str(v) for v in r[:MAX_COLS]] for r in itertools.islice(rows, MAX_ROWS + 1)]
    if len(out) > MAX_ROWS and any(c.strip() for c in out[MAX_ROWS]):
        raise ValueError(f"more than {MAX_ROWS:,} rows: split the file")
    return out[:MAX_ROWS]


def _html_rows(text):
    """Table rows with colspan/rowspan expanded into blank cells, so columns stay aligned."""
    from bs4 import BeautifulSoup
    below = {}  # column -> rows a rowspan from above still covers
    for tr in BeautifulSoup(text, "lxml").find_all("tr"):
        row = []
        for td in tr.find_all(["td", "th"], recursive=False):
            while below.get(len(row)):
                below[len(row)] -= 1
                row.append("")
            span, down = min(int(td.get("colspan") or 1), MAX_COLS), min(int(td.get("rowspan") or 1), MAX_ROWS)
            if down > 1:
                below.update(dict.fromkeys(range(len(row), len(row) + span), down - 1))
            row += [td.get_text(" ", strip=True)] + [""] * (span - 1)
            if len(row) >= MAX_COLS:
                break
        for c in below:
            if c >= len(row) and below[c]:
                below[c] -= 1
        yield row


def _kis(data):
    rows = _grid(data)
    h = next((i for i, r in enumerate(rows) if _first(r).startswith("거래일")), None)
    if h is None:
        raise ValueError(f"KIS: no header row starting with '거래일' (first rows: {rows[:3]})")
    # 2줄보기: a second header row, and every transaction spans two rows (A starts with the date).
    two = h + 1 < len(rows) and not _date(_first(rows[h + 1])) and any(c.strip() for c in rows[h + 1])
    heads, col = rows[h:h + 1 + two], {}
    for r, head in enumerate(heads):
        for c, name in enumerate(head):
            col.setdefault(_norm(name), (r, c))
    col.pop("", None)
    memo_rc = next((col[k] for k in ("거래종류", "거래상세유형") if k in col), None)
    cash_tab = "입금액" in col  # 입출금거래내역; otherwise 전체거래내역
    need = ("입금액", "출금액") if cash_tab else ("거래금액",)
    if memo_rc is None or any(k not in col for k in need):
        raise ValueError(f"KIS: expected columns 거래종류 and {'/'.join(need)}; found {list(col)}")

    mr, mc = memo_rc  # the text beside a code-only 거래종류 comes only from a column no header names (not 종목명)
    beside = not any(mc + 1 < len(head) and head[mc + 1].strip() for head in heads)
    i = h + 1 + two
    while i < len(rows):
        d = _date(_first(rows[i]))
        if not d:
            i += 1
            continue  # blank, 합계 or page junk
        # a row B never starts with a date: a one-row file taken for 2줄보기 (a junk line under its header) steps by 1
        step = 2 if two and i + 1 < len(rows) and not _date(_first(rows[i + 1])) else 1
        pair, i = rows[i:i + step], i + step

        def cell(r, c):
            return pair[r][c].strip() if r < len(pair) and c < len(pair[r]) else ""

        f = {k: cell(*rc) for k, rc in col.items()}
        memo = " ".join(next((m for m in (cell(mr, mc), cell(mr, mc + 1) if beside else "") if HANGUL.search(m)),
                             cell(mr, mc)).split())
        try:
            n = {k: _num(f.get(k, "")) for k in ("거래금액", "정산금액", "수수료", "거래세", "세금", "환율", "입금액", "출금액")}
            bal = next((_num(f[k]) for k in ("잔액", "예수금잔액", "예수금") if f.get(k)), "")
        except ValueError as e:
            raise ValueError(f"KIS row dated {d}: {e} in {pair}") from None
        kind = None
        if cash_tab:
            amount = n["입금액"] - n["출금액"]
        else:
            amount = _cash(memo, n["거래금액"], n["수수료"], n["세금"], n["거래세"])
            if _income(memo) and ("세금" not in col or (n["정산금액"] and abs(n["정산금액"] - amount) > 1)):
                kind = "UNKNOWN"  # can't confirm the income is net of withholding: human check
            if n["환율"]:
                amount *= n["환율"]
                if (kind or db.classify(memo, amount)) != "INTERNAL":
                    kind = "UNKNOWN"  # 거래금액's currency is a guess: keep it out of flows/dividends until reviewed
        yield d, amount, memo, bal, {k: v for k, v in f.items() if v}, kind


def _pdf_text(data):
    if data[:5] != b"%PDF-":
        raise ValueError("not a PDF file")
    exe = shutil.which("pdftotext")
    if not exe:
        raise RuntimeError("pdftotext not found: install poppler (apt install poppler-utils)")
    try:
        p = subprocess.run([exe, "-layout", "-enc", "UTF-8", "-", "-"], input=data, capture_output=True,
                           timeout=PDF_TIMEOUT)
    except subprocess.TimeoutExpired:
        raise ValueError(f"pdftotext ran over {PDF_TIMEOUT}s: not a 거래내역서") from None
    if p.returncode:
        raise ValueError(f"pdftotext failed: {p.stderr.decode(errors='replace').strip()}")
    if len(p.stdout) > MAX_TEXT:
        raise ValueError(f"pdftotext printed over {MAX_TEXT >> 20}MB: not a 거래내역서")
    return p.stdout.decode("utf-8")


def _toss(data):
    """정산금액 is the cash effect (net of fee/tax), in KRW on the first line even for USD rows; its sign comes from
    거래구분. Every row is checked against the running 잔액 of its currency, so a misread column fails loudly."""
    lines = _pdf_text(data).splitlines()
    tail, balance = None, {}
    for i, line in enumerate(lines):
        flat = line.replace(" ", "")
        if "원화거래내역" in flat:
            tail = KRW_TAIL
        elif "달러거래내역" in flat:
            tail = USD_TAIL
        if not TOSS_DATE.match(line):
            continue  # headers, page numbers, (ISIN) lines
        if tail is None:
            raise ValueError(f"toss: row before any 원화/달러 거래내역 section: {line.strip()!r}")
        # date, 거래구분 [종목명], [환율], then the fixed tail columns
        m = re.match(rf"\s*(\d{{4}}\.\d\d\.\d\d)\s+(\S.*?)((?:\s+{NUM}){{{len(tail)},{len(tail) + 1}}})\s*$", line)
        if not m:
            raise ValueError(f"toss: unexpected row layout: {line.strip()!r}")
        w = m[2].split()
        n = 2 if " ".join(w[:2]) in TOSS_MULTI else 1
        kind_word, product = " ".join(w[:n]), " ".join(w[n:])
        nums = [_num(x) for x in m[3].split()]
        v, raw = dict(zip(tail, nums[-len(tail):])), {"line": line.strip()}
        cash = _sign(kind_word) * v["정산금액"]
        currency, bal, change = "KRW", v["잔액"], cash
        if tail is USD_TAIL:
            block = itertools.takewhile(lambda s: not TOSS_DATE.match(s), lines[i + 1:])
            usd_line = next((s for s in block if "($" in s), "")
            usd = [_num(t) for t in USD_VALUES.findall(usd_line)]
            if len(usd) != len(USD_LINE):
                raise ValueError(f"toss: USD row without its {len(USD_LINE)} '($ …)' values: "
                                 f"{line.strip()!r} / {usd_line.strip()!r}")
            u = dict(zip(USD_LINE, usd))
            currency, bal, change = "USD", u["잔액"], _sign(kind_word) * u["정산금액"]
            raw.update(usd=change, usd_line=usd_line.strip())  # infer.py subtracts a dividend's usd from the USD cash
        if currency in balance and abs(balance[currency] + change - bal) > (0.011 if currency == "USD" else 1):
            raise ValueError(f"toss: {currency} 잔액 {balance[currency]} + {change} != {bal} at {line.strip()!r}")
        balance[currency] = bal
        memo = f"{kind_word} {product}".strip()
        kind = TOSS_KINDS.get(kind_word) or db.classify(kind_word, cash)  # never classify on the 종목명 (…배당… ETFs)
        yield date(*map(int, m[1].split("."))), cash, memo, bal, raw, kind


def _csv(data):
    rows = _grid(data)
    for h, head in enumerate(rows):
        names = [_norm(c) for c in head]
        col = {k: next((names.index(a) for a in alias if a in names), None) for k, alias in CSV_COLS.items()}
        if col["date"] is not None:
            break
    else:
        raise ValueError("csv: no header row with a date column (날짜/거래일자/date)")
    split = col["in"] is not None and col["out"] is not None
    if col["memo"] is None or not (split or col["amount"] is not None):
        raise ValueError(f"csv: need a memo column (적요/내용/memo) and 금액/amount or 입금액+출금액; header: {head}")

    def get(r, k):
        return r[col[k]].strip() if col[k] is not None and col[k] < len(r) else ""

    for r in rows[h + 1:]:
        if not any(c.strip() for c in r):
            continue
        d, memo = _date(get(r, "date")), " ".join(get(r, "memo").split())
        if not d and re.search(r"\d", get(r, "date")):
            raise ValueError(f"csv: can't read the date {get(r, 'date')!r}; use YYYY-MM-DD, YYYY.MM.DD or YYYYMMDD")
        try:
            amount = _num(get(r, "in")) - abs(_num(get(r, "out"))) if split else _num(get(r, "amount"))
            bal = _num(get(r, "bal")) if get(r, "bal") else ""
        except ValueError as e:
            raise ValueError(f"csv row {r}: {e}") from None
        raw = {head[c]: v for c, v in enumerate(r) if c < len(head) and v.strip()}
        yield d, amount if d else None, memo, bal, raw, None  # rows without a digit in the date (합계…) are skipped


PARSERS = {"kis_xlsx": _kis, "toss_pdf": _toss, "csv": _csv}
